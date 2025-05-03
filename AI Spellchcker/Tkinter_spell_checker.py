import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document
import language_tool_python
import difflib
from io import BytesIO

# Initialize language tool for grammar and spelling check
tool = language_tool_python.LanguageTool('en-US')

# Function to read docx
def read_docx(file):
    doc = Document(file)
    return '\n'.join([para.text for para in doc.paragraphs])

# Function to generate DOCX
def generate_docx(text):
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Function to check grammar and spelling
def check_text():
    text_to_check = input_text.get("1.0", "end-1c")  # Get user input text

    if not text_to_check.strip():
        messagebox.showwarning("Input Error", "Please enter or upload text!")
        return

    matches = tool.check(text_to_check)
    if not matches:
        messagebox.showinfo("Success", "No spelling or grammar mistakes found!")
        return

    corrected_text = language_tool_python.utils.correct(text_to_check, matches)
    
    # Display changes in a user-friendly way
    changes_text = ""
    for change in difflib.ndiff(text_to_check.split(), corrected_text.split()):
        if change.startswith("- "):
            wrong = change[2:]
        elif change.startswith("+ "):
            right = change[2:]
            changes_text += f"❌ **{wrong}** ➡️ ✅ **{right}**\n"

    # Display results in the Text widget
    output_text.delete("1.0", "end")
    output_text.insert(tk.END, corrected_text)

    # Show the changes in the Changes window
    changes_output.delete("1.0", "end")
    changes_output.insert(tk.END, changes_text)

# Function to upload and read .docx
def upload_file():
    file_path = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
    if file_path:
        with open(file_path, "rb") as f:
            text_to_check = read_docx(f)
            input_text.delete("1.0", "end")
            input_text.insert(tk.END, text_to_check)

# Function to save the corrected text as .docx
def save_corrected_text():
    corrected_text = output_text.get("1.0", "end-1c")
    docx_data = generate_docx(corrected_text)
    file_path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
    if file_path:
        with open(file_path, "wb") as f:
            f.write(docx_data.getvalue())
        messagebox.showinfo("Success", "Corrected text saved successfully!")

# Setup Tkinter window
window = tk.Tk()
window.title("🧠 AI Spell & Grammar Checker")
window.geometry("600x600")
window.config(bg="#F7F7F7")

# Title label with enhanced font style and color
title_label = tk.Label(window, text="AI Spell & Grammar Checker", font=("Helvetica", 16, "bold"), bg="#F7F7F7", fg="#4CAF50")
title_label.pack(pady=20)

# Instructions for the user
instructions = tk.Label(window, text="Upload a Word document or paste text to check spelling and grammar.", bg="#F7F7F7", fg="#333333")
instructions.pack(pady=10)

# Text area for user input
input_text = tk.Text(window, height=10, width=50, font=("Arial", 12), padx=10, pady=10, bd=2, relief="solid", wrap="word")
input_text.pack(pady=20)

# Buttons section
button_frame = tk.Frame(window, bg="#F7F7F7")
button_frame.pack(pady=10)

check_button = tk.Button(button_frame, text="Check Spelling & Grammar", command=check_text, font=("Arial", 12, "bold"), bg="#4CAF50", fg="white", relief="solid", padx=20, pady=10)
check_button.grid(row=0, column=0, padx=20)

upload_button = tk.Button(button_frame, text="Upload Word File", command=upload_file, font=("Arial", 12, "bold"), bg="#007BFF", fg="white", relief="solid", padx=20, pady=10)
upload_button.grid(row=0, column=1, padx=20)

save_button = tk.Button(button_frame, text="Save Corrected Text", command=save_corrected_text, font=("Arial", 12, "bold"), bg="#FF5722", fg="white", relief="solid", padx=20, pady=10)
save_button.grid(row=1, column=0, columnspan=2, pady=10)

# Output area to display corrected text
output_text = tk.Text(window, height=10, width=50, font=("Arial", 12), padx=10, pady=10, bd=2, relief="solid", wrap="word")
output_text.pack(pady=20)

# Changes display window
changes_label = tk.Label(window, text="Corrections Made", font=("Helvetica", 14, "bold"), bg="#F7F7F7", fg="#4CAF50")
changes_label.pack(pady=10)

changes_output = tk.Text(window, height=5, width=50, font=("Arial", 12), padx=10, pady=10, bd=2, relief="solid", wrap="word")
changes_output.pack(pady=10)

# Run the app
window.mainloop()