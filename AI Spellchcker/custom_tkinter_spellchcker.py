import customtkinter as ctk
from tkinter import messagebox, filedialog
from docx import Document
import fitz  # PyMuPDF
from PIL import Image
import os
import language_tool_python
import time

# Initialize grammar and spell checker
tool = language_tool_python.LanguageTool('en-US')

# Correct text using language_tool_python with progress
def correct_text(text, progress_bar):
    matches = tool.check(text)
    total = len(matches) if matches else 1
    corrected_text = language_tool_python.utils.correct(text, matches)

    # Fake progress bar loop
    for i in range(10):
        time.sleep(0.05)  # Simulate work
        progress_bar.set((i + 1) / 10)

    return corrected_text

# Correct paragraph text from UI
def correct_paragraph_text(input_box, output_box, progress_bar):
    input_text = input_box.get("1.0", ctk.END)
    corrected = correct_text(input_text, progress_bar)
    output_box.delete("1.0", ctk.END)
    output_box.insert(ctk.END, corrected)
    progress_bar.set(0)

# Upload DOCX file and correct
def upload_docx(progress_bar):
    filepath = filedialog.askopenfilename(filetypes=[("Word files", "*.docx")])
    if not filepath:
        return
    doc = Document(filepath)
    full_text = '\n'.join([para.text for para in doc.paragraphs])
    corrected_text = correct_text(full_text, progress_bar)

    new_doc = Document()
    new_doc.add_paragraph(corrected_text)
    output_path = "corrected_" + os.path.basename(filepath)
    new_doc.save(output_path)
    messagebox.showinfo("Success", f"Corrected DOCX saved as:\n{output_path}")
    progress_bar.set(0)

# Upload PDF file and correct
def upload_pdf(progress_bar):
    filepath = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
    if not filepath:
        return
    doc = fitz.open(filepath)
    full_text = ""
    for page in doc:
        full_text += page.get_text()
    corrected_text = correct_text(full_text, progress_bar)

    output_path = "corrected_" + os.path.basename(filepath)
    new_pdf = fitz.open()
    page = new_pdf.new_page()
    page.insert_text((72, 72), corrected_text, fontsize=11)
    new_pdf.save(output_path)
    messagebox.showinfo("Success", f"Corrected PDF saved as:\n{output_path}")
    progress_bar.set(0)

# Show the spell checker window after login
def show_spell_checker():
    global input_box, output_box
    app.destroy()

    spell_checker_app = ctk.CTk()
    spell_checker_app.title("AI Grammar & Spell Checker")
    spell_checker_app.geometry("800x800")

    title_font = ("ZORGEOUS", 28, "bold")
    label_font = ("ZORGEOUS", 16, "bold")

    main_frame = ctk.CTkFrame(spell_checker_app, corner_radius=10, fg_color="#FFFFFF")
    main_frame.pack(padx=20, pady=20, fill="both", expand=True)

    title_label = ctk.CTkLabel(main_frame, text="🧠 AI Grammar & Spell Checker", font=title_font, text_color="#000000")
    title_label.pack(pady=(10, 5))

    desc_label = ctk.CTkLabel(main_frame, text="Check grammar in text, DOCX or PDF files.", font=label_font, text_color="#000000")
    desc_label.pack(pady=(0, 20))

    input_label = ctk.CTkLabel(main_frame, text="✍️ Enter or Paste Text Below", font=label_font, text_color="#000000")
    input_label.pack(anchor="w", padx=20)

    input_box = ctk.CTkTextbox(main_frame, height=200, font=("ZORGEOUS", 14), wrap="word", fg_color="#E0E0E0", border_width=2, border_color="#B0B0B0")
    input_box.pack(padx=20, pady=10, fill="both")

    correct_button = ctk.CTkButton(main_frame, text="✅ Correct Paragraph Text", command=lambda: correct_paragraph_text(input_box, output_box, progress_bar), fg_color="#4CAF50", hover_color="#388E3C")
    correct_button.pack(pady=10)

    output_label = ctk.CTkLabel(main_frame, text="🧾 Corrected Output", font=label_font, text_color="#000000")
    output_label.pack(anchor="w", padx=20)

    output_box = ctk.CTkTextbox(main_frame, height=200, font=("ZORGEOUS", 14), wrap="word", fg_color="#E0E0E0", border_width=2, border_color="#B0B0B0")
    output_box.pack(padx=20, pady=10, fill="both")

    progress_bar = ctk.CTkProgressBar(main_frame, width=400)
    progress_bar.pack(pady=10)
    progress_bar.set(0)

    docx_button = ctk.CTkButton(main_frame, text="📄 Upload DOCX File", command=lambda: upload_docx(progress_bar), fg_color="#2196F3", hover_color="#1976D2")
    docx_button.pack(pady=(10, 5))

    pdf_button = ctk.CTkButton(main_frame, text="📄 Upload PDF File", command=lambda: upload_pdf(progress_bar), fg_color="#2196F3", hover_color="#1976D2")
    pdf_button.pack(pady=(0, 20))

    spell_checker_app.mainloop()

# Login window
app = ctk.CTk()
app.geometry("800x500")
app.title("Login")
ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")

frame = ctk.CTkFrame(app, fg_color="#FFFFFF")
frame.pack(fill="both", expand=True)

left_panel = ctk.CTkFrame(frame, width=300, fg_color="#F0F0F0")
left_panel.pack(side="left", fill="both")
image = ctk.CTkImage(light_image=Image.open("login_bg.png"), size=(300, 500))
image_label = ctk.CTkLabel(left_panel, image=image, text="")
image_label.pack(fill="both", expand=True)

right_panel = ctk.CTkFrame(frame, fg_color="#FFFFFF")
right_panel.pack(side="right", fill="both", expand=True, padx=60, pady=60)

title = ctk.CTkLabel(right_panel, text="Welcome Back!", font=("ZORGEOUS", 28, "bold"), text_color="#6A0DAD")
title.pack(anchor="w", pady=(0, 10))

subtitle = ctk.CTkLabel(right_panel, text="Sign in to your account", font=("ZORGEOUS", 16, "bold"))
subtitle.pack(anchor="w", pady=(0, 30))

email_entry = ctk.CTkEntry(right_panel, width=300, placeholder_text="Email", font=("ZORGEOUS", 14))
email_entry.pack(pady=10)

password_entry = ctk.CTkEntry(right_panel, width=300, placeholder_text="Password", show="*", font=("ZORGEOUS", 14))
password_entry.pack(pady=10)

def login_action():
    email = email_entry.get()
    password = password_entry.get()
    if email == "admin" and password == "admin":
        messagebox.showinfo("Login Success", "Welcome!")
        show_spell_checker()
    else:
        messagebox.showerror("Error", "Invalid credentials")

login_btn = ctk.CTkButton(right_panel, text="Login", command=login_action, fg_color="#6A0DAD", hover_color="#4B0082")
login_btn.pack(pady=20)

app.mainloop()