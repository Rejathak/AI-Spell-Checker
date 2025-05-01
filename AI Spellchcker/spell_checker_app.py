import streamlit as st
from docx import Document
import language_tool_python
import difflib
from io import BytesIO
import PyPDF2

# App title
st.title("🧠 AI Spell & Grammar Checker")
st.write("Check spelling and grammar from typed text, Word docs, or PDF files.")

tool = language_tool_python.LanguageTool('en-US')

uploaded_docx = st.file_uploader("📎 Upload a Word (.docx) file", type=["docx"])
uploaded_pdf = st.file_uploader("📄 Upload a PDF file", type=["pdf"])
user_input = st.text_area("Or enter/paste your text here:")

text_to_check = ""

# DOCX extraction
if uploaded_docx:
    def read_docx(file):
        doc = Document(file)
        return '\n'.join([para.text for para in doc.paragraphs])
    text_to_check = read_docx(uploaded_docx)
    st.success("✅ Text extracted from Word file.")

# PDF extraction
elif uploaded_pdf:
    reader = PyPDF2.PdfReader(uploaded_pdf)
    pdf_text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pdf_text += page_text
    if pdf_text.strip():
        text_to_check = pdf_text
        st.success("✅ Text extracted from PDF file.")
    else:
        st.warning("⚠️ Could not extract text from the PDF. Is it a scanned image?")

elif user_input.strip():
    text_to_check = user_input

# Function to generate .docx
def generate_docx(text):
    doc = Document()
    doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Check button
if st.button("Check Spelling & Grammar") and text_to_check.strip():
    matches = tool.check(text_to_check)

    if not matches:
        st.success("🎉 No spelling or grammar mistakes found!")
    else:
        corrected_text = language_tool_python.utils.correct(text_to_check, matches)

        st.error(f"❌ Found {len(matches)} issue(s):")
        st.markdown("### 🛠️ Corrections Made:")
        diff = difflib.ndiff(text_to_check.split(), corrected_text.split())
        for change in diff:
            if change.startswith("- "):
                wrong = change[2:]
            elif change.startswith("+ "):
                right = change[2:]
                st.markdown(f"- ❌ **{wrong}** ➡️ ✅ **{right}**")

        st.markdown("### 📝 Original Text")
        st.text_area("Original:", text_to_check, height=150)

        st.markdown("### ✅ Corrected Text")
        st.text_area("Corrected:", corrected_text, height=150)

        # DOCX download
        if uploaded_docx or user_input or uploaded_pdf:
            docx_data = generate_docx(corrected_text)
            st.download_button(
                label="📄 Download as Word (.docx)",
                data=docx_data,
                file_name="corrected_output.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    